"""AI 菜單：正式草稿、換菜/鎖定/分數、公版 Excel，以及完整練習沙盒。"""
from __future__ import annotations
import json,os,re,secrets
from collections import defaultdict
from datetime import date,datetime,timedelta
from io import BytesIO
from types import SimpleNamespace
from flask import abort,current_app,flash,redirect,render_template,request,send_file,session,url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font
from sqlalchemy.orm import joinedload
from blueprints.order_tool import _commit,_date,_decimal,_int,order_bp
from extensions import db
from models import KitchenIngredient,KitchenMenuPlan,KitchenMenuPlanItem,KitchenRecipe,KitchenRecipeIngredient
from ai_models import KitchenMealProfile,KitchenMenuDraft,KitchenMenuDraftItem
from services import meal_profiles as profile_service, recipe_tags as tag_service
from services.ai_menu_engine import CATEGORY_ORDER,DEFAULT_STRUCTURE,STRUCTURE_LIMITS,RuleFeasibilityError,Rules,build_structure,candidate_allowed,describe_rules,generate,hard_rule_violations,rebuild,replacement_options,service_dates,variant_violations,WEEKDAY_LABEL
from services.nutrition import recipe_nutrition
from services.nutrition_sources import STATUS_ESTIMATED,STATUS_MATCHED,match_ingredient_name
from services.practice_database import ensure_practice_database
WEEKDAY="一二三四五六日"
PRACTICE_USER="practice"

def _practice_ok():
    return bool(session.get("kitchen_practice"))

def _default_range():
    today=date.today(); start=(today.replace(day=1)+timedelta(days=32)).replace(day=1); end=(start+timedelta(days=32)).replace(day=1)-timedelta(days=1); return start,end

def _seed_profiles():
    if profile_service.seed_profiles(db.session):db.session.commit()

def _readonly_profiles():
    rows=profile_service.active_profiles()
    if rows:return rows
    return [SimpleNamespace(id=code,code=code,name=name,kcal_min=kmin,kcal_max=kmax,source_document=doc,source_version=ver,note=note,kcal_label=f"{kmin}～{kmax} kcal") for code,name,kmin,kmax,_,doc,ver,_,note in profile_service.PROFILE_SEEDS]

def _practice_profile(raw):
    for p in _readonly_profiles():
        if str(p.id)==str(raw) or p.code==raw:return p

def _parse_rules(form):
    return Rules.from_dict({"recipe_repeat_days":form.get("recipe_repeat_days"),"main_repeat_days":form.get("main_repeat_days"),"fish_per_week_min":form.get("fish_per_week_min"),"fish_per_week_max":form.get("fish_per_week_max"),"fried_per_week_max":form.get("fried_per_week_max"),"sweet_soup_per_week_max":form.get("sweet_soup_per_week_max"),"fish_weekdays":form.getlist("fish_weekdays"),"fried_weekdays":form.getlist("fried_weekdays"),"sweet_soup_weekdays":form.getlist("sweet_soup_weekdays"),"prefer_kcal_in_range":form.get("prefer_kcal_in_range")=="1","exclude_incomplete_nutrition":form.get("exclude_incomplete_nutrition")=="1"})

def _recipe_ingredients(ids):
    if not ids:return {}
    rows=(KitchenRecipe.query.options(joinedload(KitchenRecipe.ingredients).joinedload(KitchenRecipeIngredient.ingredient)).filter(KitchenRecipe.id.in_(ids)).all()); out={}
    for r in rows:
        parts=[]
        for x in r.ingredients or ():
            if not x.ingredient:continue
            try:a=f"{float(x.grams_per_person):g}" if x.grams_per_person is not None else ""
            except Exception:a=str(x.grams_per_person or "")
            parts.append(f"{x.ingredient.name}{(' '+a+(x.ingredient.base_unit or 'g')) if a else ''}")
        out[r.id]="、".join(parts) if parts else "尚未建立配方"
    return out

def _view_rows(result,draft=None):
    ing=_recipe_ingredients({rid for rid in result.assignment.values() if rid}); locks={}
    if draft:locks={(x.service_date,x.category,x.slot_index):x.locked for x in draft.items}
    rows=[]
    for i,d in enumerate(result.dates):
        cells=[]
        for c in CATEGORY_ORDER:
            for s in range(1,result.structure.get(c,0)+1):
                cand=result.recipe_at(i,c,s); cells.append({"category":c,"slot_index":s,"recipe_id":cand.id if cand else None,"name":cand.name if cand else "","kcal":cand.kcal if cand else None,"ingredients":ing.get(cand.id,"") if cand else "","tags":sorted(cand.tags) if cand else [],"locked":locks.get((d,c,s),False)})
        rep=result.days[d]; rows.append({"service_date":d,"weekday":WEEKDAY[d.weekday()],"cells":cells,"kcal":rep.kcal,"warnings":rep.warnings,"ok_messages":rep.ok_messages,"day_locked":bool(cells) and all(x["locked"] for x in cells)})
    return rows

def _headers(structure):
    return [c if structure[c]==1 else f"{c}{s}" for c in CATEGORY_ORDER for s in range(1,structure.get(c,0)+1)]

def _draft_result(draft):
    metadata=json.loads(draft.rules_json or "{}");variant=metadata.get("_meal_variant","legacy")
    structure=build_structure(json.loads(draft.structure_json or "{}")); rules=Rules.from_dict(metadata); dates=service_dates(draft.start_date,draft.end_date,draft.include_weekends); index={d:i for i,d in enumerate(dates)}; assignment={}
    for x in draft.items:
        if x.service_date in index:assignment[(index[x.service_date],x.category,x.slot_index)]=x.recipe_id
    p=draft.profile; return rebuild(dates,structure,rules,assignment,int(p.kcal_min) if p and p.kcal_min is not None else None,int(p.kcal_max) if p and p.kcal_max is not None else None,variant)

def _draft_metadata(draft):
    try:return json.loads(draft.rules_json or "{}")
    except (TypeError,ValueError):return {}

def _paired_draft(draft):
    pair_key=_draft_metadata(draft).get("_pair_key")
    if not pair_key:return None
    for row in KitchenMenuDraft.query.filter(KitchenMenuDraft.id!=draft.id).order_by(KitchenMenuDraft.id.desc()).limit(50):
        if _draft_metadata(row).get("_pair_key")==pair_key:return row
    return None

def _save_result(draft,result,keep_locked=False):
    existing={(x.service_date,x.category,x.slot_index):x for x in draft.items}; wanted=set()
    for key,rid in result.assignment.items():
        i,c,s=key; d=result.dates[i]; k=(d,c,s); wanted.add(k); row=existing.get(k)
        if row and keep_locked and row.locked:continue
        if not row:row=KitchenMenuDraftItem(draft_id=draft.id,service_date=d,category=c,slot_index=s);db.session.add(row)
        row.recipe_id=rid;row.sort_order=CATEGORY_ORDER.index(c)*10+s
    for k,row in existing.items():
        if k not in wanted and not (keep_locked and row.locked):db.session.delete(row)

def _settings(form,profile):
    start=_date(form.get("start_date"));end=_date(form.get("end_date"))
    if not profile:return None,"請先選擇供餐對象。"
    if not start or not end or end<start:return None,"請填寫正確的排菜起訖日期。"
    if (end-start).days>92:return None,"一次最多排 3 個月。"
    structure=build_structure({c:form.get(f"structure_{c}") for c in CATEGORY_ORDER})
    if not sum(structure.values()):return None,"每日至少要有一道菜。"
    return (start,end,structure,_parse_rules(form),form.get("include_weekends")=="1"),None

@order_bp.get("/ai-menu")
def ai_menu():
    _seed_profiles();    start,end=_default_range();drafts=KitchenMenuDraft.query.options(joinedload(KitchenMenuDraft.profile)).order_by(KitchenMenuDraft.id.desc()).limit(20).all()
    default_rules=Rules()
    cap_rows=(
        ("fish","魚類","fish_per_week_max","fish_weekdays",default_rules.fish_per_week_max),
        ("fried","炸物","fried_per_week_max","fried_weekdays",default_rules.fried_per_week_max),
        ("sweet_soup","甜湯","sweet_soup_per_week_max","sweet_soup_weekdays",default_rules.sweet_soup_per_week_max),
    )
    return render_template("kitchen/ai_menu.html",profiles=profile_service.active_profiles(),drafts=drafts,categories=CATEGORY_ORDER,default_structure=DEFAULT_STRUCTURE,structure_limits=STRUCTURE_LIMITS,default_rules=default_rules,weekday_options=list(enumerate(WEEKDAY_LABEL)),cap_rows=cap_rows,start_default=start,end_default=end,tagged_recipe_count=0,practice=_practice_ok())

@order_bp.post("/ai-menu/generate")
def ai_menu_generate():
    pid=_int(request.form.get("profile_id"),default=0);profile=db.session.get(KitchenMealProfile,pid) if pid else None; settings,error=_settings(request.form,profile)
    if error:flash(error,"error");return redirect(url_for("order_tool.ai_menu"))
    start,end,structure,rules,weekends=settings
    vegetarian_rules=Rules.from_dict(rules.to_dict());vegetarian_rules.fish_per_week_min=0;vegetarian_rules.fish_weekdays=()
    try:
        regular_result=generate(start,end,structure,rules,int(profile.kcal_min),int(profile.kcal_max),weekends,meal_variant="regular")
        vegetarian_result=generate(start,end,structure,vegetarian_rules,int(profile.kcal_min),int(profile.kcal_max),weekends,meal_variant="vegetarian",reference_assignment=regular_result.assignment)
    except RuleFeasibilityError as exc:
        flash(f"無法產生符合硬限制的菜單：{exc}","error");return redirect(url_for("order_tool.ai_menu"))
    base_name=(request.form.get("name") or "").strip() or f"{profile.name} {start:%Y/%m}";pair_key=secrets.token_hex(12);drafts=[]
    for variant,label,variant_rules,result in (("regular","葷食",rules,regular_result),("vegetarian","素食",vegetarian_rules,vegetarian_result)):
        rule_data=variant_rules.to_dict();rule_data.update({"_meal_variant":variant,"_pair_key":pair_key})
        draft=KitchenMenuDraft(name=f"{base_name}｜{label}",profile_id=profile.id,start_date=start,end_date=end,include_weekends=weekends,meal_type="午餐",structure_json=json.dumps(structure,ensure_ascii=False),rules_json=json.dumps(rule_data,ensure_ascii=False),status="draft");db.session.add(draft);db.session.flush();_save_result(draft,result);drafts.append(draft)
    db.session.commit();return redirect(url_for("order_tool.ai_menu_draft",draft_id=drafts[0].id))

@order_bp.get("/ai-menu/drafts/<int:draft_id>")
def ai_menu_draft(draft_id):
    draft=db.session.get(KitchenMenuDraft,draft_id)
    if not draft:abort(404)
    result=_draft_result(draft);paired=_paired_draft(draft);return render_template("kitchen/ai_menu_result.html",draft=draft,result=result,rows=_view_rows(result,draft),headers=_headers(result.structure),rule_lines=describe_rules(result.rules),tag_label=tag_service.tag_label,practice=_practice_ok(),meal_variant=result.meal_variant,paired_draft=paired,paired_variant=_draft_metadata(paired).get("_meal_variant") if paired else None)

@order_bp.post("/ai-menu/drafts/<int:draft_id>/regenerate")
def ai_menu_regenerate(draft_id):
    draft=db.session.get(KitchenMenuDraft,draft_id)
    if not draft:abort(404)
    old=_draft_result(draft);locked={}
    if request.form.get("keep_locked")=="1":
        index={d:i for i,d in enumerate(old.dates)}
        for x in draft.items:
            if x.locked and x.service_date in index:locked[(index[x.service_date],x.category,x.slot_index)]=x.recipe_id
    p=draft.profile
    paired=_paired_draft(draft);reference=_draft_result(paired).assignment if old.meal_variant=="vegetarian" and paired else None
    try:new=generate(draft.start_date,draft.end_date,old.structure,old.rules,int(p.kcal_min),int(p.kcal_max),draft.include_weekends,locked=locked,meal_variant=old.meal_variant,reference_assignment=reference)
    except RuleFeasibilityError as exc:
        flash(f"無法重新排菜：{exc}","error");return redirect(url_for("order_tool.ai_menu_draft",draft_id=draft.id))
    _save_result(draft,new,keep_locked=True);return _commit("已重新排菜；鎖定菜色保持不變。","order_tool.ai_menu_draft",draft_id=draft.id)

@order_bp.get("/ai-menu/drafts/<int:draft_id>/swap")
def ai_menu_swap(draft_id):
    draft=db.session.get(KitchenMenuDraft,draft_id)
    if not draft:abort(404)
    d=_date(request.args.get("date"));cat=(request.args.get("category") or "").strip();slot=_int(request.args.get("slot"),default=1);result=_draft_result(draft)
    if d not in result.dates or cat not in CATEGORY_ORDER:abort(404)
    key=(result.dates.index(d),cat,slot);options=replacement_options(result,key)
    ingredients=KitchenIngredient.query.filter_by(active=True).order_by(KitchenIngredient.name).all()
    return render_template("kitchen/ai_menu_swap.html",draft=draft,result=result,service_date=d,weekday=WEEKDAY[d.weekday()],category=cat,slot_index=slot,options=options,current_id=result.assignment.get(key),day_report=result.days[d],recipe_options=[{"id":candidate.id,"name":candidate.name,"category":candidate.category} for candidate,_penalty,_current in options],ingredient_options=[{"id":ingredient.id,"name":ingredient.name,"base_unit":ingredient.base_unit or "g","purchase_unit":ingredient.purchase_unit} for ingredient in ingredients])

@order_bp.post("/ai-menu/drafts/<int:draft_id>/swap/new")
def ai_menu_swap_new(draft_id):
    draft=db.session.get(KitchenMenuDraft,draft_id)
    if not draft:abort(404)
    d=_date(request.form.get("date"));cat=(request.form.get("category") or "").strip();slot=_int(request.form.get("slot"),default=1)
    name=(request.form.get("name") or "").strip()
    result=_draft_result(draft)
    if d not in result.dates or cat not in CATEGORY_ORDER:abort(404)
    redirect_args={"draft_id":draft.id,"date":d.isoformat(),"category":cat,"slot":slot}
    if not name:
        flash("請輸入新菜色名稱。","error");return redirect(url_for("order_tool.ai_menu_swap",**redirect_args))
    if KitchenRecipe.query.filter_by(name=name).first():
        flash("菜色名稱已存在，請直接在上方搜尋。","error");return redirect(url_for("order_tool.ai_menu_swap",**redirect_args))
    components={}
    for raw_id,raw_amount in zip(request.form.getlist("ingredient_id"),request.form.getlist("grams_per_person")):
        ingredient_id=_int(raw_id,default=0);amount=_decimal(raw_amount,default=None)
        ingredient=db.session.get(KitchenIngredient,ingredient_id) if ingredient_id else None
        if ingredient and ingredient.active and amount is not None and amount>0:components[ingredient.id]=amount
    if not components:
        flash("至少要選一項食材並填寫正確的每人用量。","error");return redirect(url_for("order_tool.ai_menu_swap",**redirect_args))
    # Temporarily include the pending recipe so the normal AI hard-rule validator
    # can evaluate the requested swap. It is disabled again before commit; new
    # recipes never join future AI candidate pools without an explicit enable.
    recipe=KitchenRecipe(name=name,category=cat,active=True,note="從 AI 換菜頁新增")
    db.session.add(recipe);db.session.flush()
    for ingredient_id,amount in components.items():
        db.session.add(KitchenRecipeIngredient(recipe_id=recipe.id,ingredient_id=ingredient_id,grams_per_person=amount,quantity_status="manual",source_note="人工確認"))
    db.session.flush()
    updated=_draft_result(draft);key=(updated.dates.index(d),cat,slot);trial=dict(updated.assignment);trial[key]=recipe.id
    candidate=updated.candidates.get(recipe.id)
    violations=variant_violations(trial,updated.candidates,updated.meal_variant)+hard_rule_violations(updated.dates,updated.rules,trial,updated.candidates)
    if not candidate_allowed(candidate,updated.meal_variant) or violations:
        db.session.rollback();flash("這道菜不符合目前葷／素分類或整週硬限制，因此沒有新增。","error");return redirect(url_for("order_tool.ai_menu_swap",**redirect_args))
    item=next((x for x in draft.items if x.service_date==d and x.category==cat and x.slot_index==slot),None)
    if not item:db.session.rollback();abort(404)
    item.recipe_id=recipe.id;recipe.active=False;db.session.commit();flash("新菜色與配方已建立，並已換入菜單。","success");return redirect(url_for("order_tool.ai_menu_draft",draft_id=draft.id))

@order_bp.post("/ai-menu/drafts/<int:draft_id>/swap")
def ai_menu_swap_apply(draft_id):
    draft=db.session.get(KitchenMenuDraft,draft_id)
    if not draft:abort(404)
    d=_date(request.form.get("date"));cat=(request.form.get("category") or "").strip();slot=_int(request.form.get("slot"),default=1);rid=_int(request.form.get("recipe_id"),default=0) or None
    item=next((x for x in draft.items if x.service_date==d and x.category==cat and x.slot_index==slot),None)
    if not item:abort(404)
    if rid and any(x.service_date==d and x.recipe_id==rid and x.id!=item.id for x in draft.items):flash("同一天不能重複同一道菜。","error");return redirect(url_for("order_tool.ai_menu_draft",draft_id=draft.id))
    result=_draft_result(draft)
    if d not in result.dates:abort(404)
    key=(result.dates.index(d),cat,slot);trial=dict(result.assignment);trial[key]=rid
    candidate=result.candidates.get(rid) if rid else None
    if rid and not candidate_allowed(candidate,result.meal_variant):
        flash("不能換成這道菜：不符合目前葷／素菜單分類。","error");return redirect(url_for("order_tool.ai_menu_draft",draft_id=draft.id))
    violations=variant_violations(trial,result.candidates,result.meal_variant)+hard_rule_violations(result.dates,result.rules,trial,result.candidates)
    if violations:
        flash("不能換成這道菜："+"；".join(violations),"error");return redirect(url_for("order_tool.ai_menu_draft",draft_id=draft.id))
    item.recipe_id=rid;return _commit("已換菜，熱量與規則分數已重新計算。","order_tool.ai_menu_draft",draft_id=draft.id)

@order_bp.post("/ai-menu/drafts/<int:draft_id>/lock-item")
def ai_menu_lock_item(draft_id):
    draft=db.session.get(KitchenMenuDraft,draft_id);d=_date(request.form.get("date"));cat=request.form.get("category");slot=_int(request.form.get("slot"),default=1);item=next((x for x in draft.items if x.service_date==d and x.category==cat and x.slot_index==slot),None)
    if not item:abort(404)
    item.locked=not item.locked;return _commit("鎖定狀態已更新。","order_tool.ai_menu_draft",draft_id=draft_id)

@order_bp.post("/ai-menu/drafts/<int:draft_id>/lock-day")
def ai_menu_lock_day(draft_id):
    draft=db.session.get(KitchenMenuDraft,draft_id);d=_date(request.form.get("date"));items=[x for x in draft.items if x.service_date==d]
    if not items:abort(404)
    target=not all(x.locked for x in items)
    for x in items:x.locked=target
    return _commit("整日鎖定狀態已更新。","order_tool.ai_menu_draft",draft_id=draft_id)

@order_bp.post("/ai-menu/drafts/<int:draft_id>/apply")
def ai_menu_apply(draft_id):
    if request.form.get("confirm_apply")!="YES":flash("未套用：請先確認。","error");return redirect(url_for("order_tool.ai_menu_draft",draft_id=draft_id))
    draft=db.session.get(KitchenMenuDraft,draft_id)
    if not draft:abort(404)
    result=_draft_result(draft);violations=variant_violations(result.assignment,result.candidates,result.meal_variant)+hard_rule_violations(result.dates,result.rules,result.assignment,result.candidates)
    if violations:
        flash("無法套用：目前菜單違反硬限制："+"；".join(violations),"error");return redirect(url_for("order_tool.ai_menu_draft",draft_id=draft.id))
    grouped=defaultdict(list)
    for x in draft.items:
        if x.recipe_id:grouped[x.service_date].append(x)
    for d,items in grouped.items():
        plan_name="中央素食菜單" if result.meal_variant=="vegetarian" else "中央菜單"
        plan=KitchenMenuPlan.query.filter_by(service_date=d,meal_type=draft.meal_type,name=plan_name).one_or_none()
        if not plan:plan=KitchenMenuPlan(service_date=d,meal_type=draft.meal_type,name=plan_name,status="draft");db.session.add(plan);db.session.flush()
        KitchenMenuPlanItem.query.filter_by(plan_id=plan.id).delete(synchronize_session=False)
        for i,x in enumerate(sorted(items,key=lambda z:z.sort_order),1):db.session.add(KitchenMenuPlanItem(plan_id=plan.id,recipe_id=x.recipe_id,sort_order=i))
    draft.status="applied";draft.applied_at=datetime.utcnow()
    return _commit(f"已套用 {len(grouped)} 個供餐日到練習中央菜單；正式資料不受影響。" if _practice_ok() else f"已套用 {len(grouped)} 個供餐日到正式中央菜單。","order_tool.ai_menu_draft",draft_id=draft.id)

@order_bp.post("/ai-menu/drafts/<int:draft_id>/delete")
def ai_menu_delete(draft_id):
    draft=db.session.get(KitchenMenuDraft,draft_id)
    if draft:db.session.delete(draft)
    return _commit("草稿已刪除。","order_tool.ai_menu")

def _safe_name(raw,fallback):
    text=re.sub(r'[\\/:*?"<>|\x00-\x1f]+',"_",(raw or "").strip()).strip(" ._")[:80] or fallback
    return text if text.lower().endswith(".xlsx") else text+".xlsx"

@order_bp.get("/ai-menu/drafts/<int:draft_id>/public.xlsx")
def ai_menu_public_excel(draft_id):
    draft=db.session.get(KitchenMenuDraft,draft_id)
    if not draft:abort(404)
    result=_draft_result(draft);rows=_view_rows(result,draft);wb=Workbook();ws=wb.active;ws.title="公版菜單";ws.append(["日期","星期","類別","菜色","食材","菜色熱量(kcal/人)","當日總熱量(kcal/人)","狀態"])
    for c in ws[1]:c.font=Font(bold=True)
    for row in rows:
        status="；".join(row["warnings"] or row["ok_messages"]) or "—"
        for cell in row["cells"]:ws.append([row["service_date"],f"週{row['weekday']}",cell["category"],cell["name"],cell["ingredients"],cell["kcal"] if cell["kcal"] is not None else "",row["kcal"] if row["kcal"] is not None else "",status])
    for col,w in {"A":14,"B":9,"C":10,"D":24,"E":48,"F":20,"G":22,"H":42}.items():ws.column_dimensions[col].width=w
    for rr in ws.iter_rows():
        for c in rr:c.alignment=Alignment(vertical="top",wrap_text=True)
    out=BytesIO();wb.save(out);out.seek(0);return send_file(out,as_attachment=True,download_name=_safe_name(request.args.get("filename"),f"公版菜單_{draft.start_date}_{draft.end_date}"),mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@order_bp.route("/ai-menu/practice/login",methods=["GET","POST"])
def ai_menu_practice_login():
    error=""
    if request.method=="POST":
        try:
            ensure_practice_database()
        except Exception:
            current_app.logger.exception("Failed to initialize kitchen practice database")
            error="練習資料庫初始化失敗，請確認 PRACTICE_DATABASE_URL 或部署儲存空間設定。"
        else:
            db.session.remove();session["kitchen_practice"]=True
            return redirect(url_for("order_tool.index"))
    return render_template("kitchen/ai_menu_practice_login.html",error=error)

@order_bp.get("/ai-menu/practice/logout")
def ai_menu_practice_logout():
    db.session.remove();session.pop("kitchen_practice",None)
    return redirect(url_for("order_tool.ai_menu_practice_login"))

@order_bp.get("/ai-menu/practice")
def ai_menu_practice():
    if not _practice_ok():return redirect(url_for("order_tool.ai_menu_practice_login"))
    return redirect(url_for("order_tool.index"))

@order_bp.post("/ai-menu/practice/generate")
def ai_menu_practice_generate():
    if not _practice_ok():return redirect(url_for("order_tool.ai_menu_practice_login"))
    return ai_menu_generate()

def _nutrition_view(ing):
    if ing.kcal_per_100g is not None:return {"kcal":ing.kcal_per_100g,"source":ing.nutrition_source or "人工輸入","food_name":ing.nutrition_food_name or "","code":ing.nutrition_code or "","status":"人工確認" if ing.nutrition_verified else "已設定","note":ing.nutrition_note or ""}
    m=match_ingredient_name(ing.name);return {"kcal":m.kcal_per_100g,"source":m.source_label,"food_name":m.food.name if m.food else ("團膳預設代表值" if m.status==STATUS_ESTIMATED else ""),"code":m.food.code if m.food else "","status":"TFDA 對應" if m.status==STATUS_MATCHED else ("估算值" if m.status==STATUS_ESTIMATED else "待確認"),"note":m.note}

@order_bp.app_template_global(name="ingredient_nutrition")
def ingredient_nutrition(ing):return _nutrition_view(ing)

@order_bp.post("/ingredients/<int:ingredient_id>/nutrition")
def ingredient_nutrition_update(ingredient_id):
    ing=db.session.get(KitchenIngredient,ingredient_id)
    if not ing:abort(404)
    raw=(request.form.get("kcal_per_100g") or "").strip()
    try:kcal=None if not raw else float(raw)
    except ValueError:flash("熱量格式不正確。","error");return redirect(url_for("order_tool.ingredients",edit=ingredient_id))
    ing.kcal_per_100g=kcal;ing.nutrition_source=(request.form.get("nutrition_source") or "").strip() or None;ing.nutrition_food_name=(request.form.get("nutrition_food_name") or "").strip() or None;ing.nutrition_note=(request.form.get("nutrition_note") or "").strip() or None;ing.nutrition_verified=request.form.get("nutrition_verified")=="1";ing.nutrition_status="verified" if ing.nutrition_verified else ("matched" if kcal is not None else "pending");db.session.commit();flash("食材熱量資料已更新。","success");return redirect(url_for("order_tool.ingredients",edit=ingredient_id))

@order_bp.get("/recipe-tags")
def recipe_tags():
    recipes=(KitchenRecipe.query.options(joinedload(KitchenRecipe.ingredients).joinedload(KitchenRecipeIngredient.ingredient),joinedload(KitchenRecipe.tags)).order_by(KitchenRecipe.category,KitchenRecipe.name).all());rows=[{"recipe":r,"manual":{x.tag for x in r.tags if x.source=="manual"},"suggested":tag_service.suggest_tags(r),"nutrition":recipe_nutrition(r)} for r in recipes];return render_template("kitchen/recipe_tags.html",rows=rows[:200],total=len(rows),tag_defs=tag_service.TAG_DEFS,categories=CATEGORY_ORDER,category="",q="",only="")

@order_bp.post("/recipe-tags")
def recipe_tags_save():
    ids=[_int(x,default=0) for x in request.form.getlist("recipe_id")];ids=[x for x in ids if x]
    if not ids:
        one=_int(request.form.get("recipe_id"),default=0);ids=[one] if one else []
    recipes=KitchenRecipe.query.options(joinedload(KitchenRecipe.tags)).filter(KitchenRecipe.id.in_(ids)).all() if ids else []
    for r in recipes:tag_service.set_manual_tags(db.session,r,set(request.form.getlist(f"tags_{r.id}")) or set(request.form.getlist("tags")))
    return _commit("菜色標記已儲存。","order_tool.recipe_tags")
