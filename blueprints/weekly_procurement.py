"""Weekly procurement review, traceability checks, and vendor export.

This module is additive: confirmed daily purchase snapshots remain untouched.
The weekly view recomputes expected demand from the current menu/BOM to flag
stale or missing purchase rows, while exports use the saved purchase rows.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
import re

from flask import Blueprint, abort, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import selectinload

from extensions import db
from models import (
    KitchenMenuAssignment,
    KitchenMenuPlan,
    KitchenMenuPlanItem,
    KitchenPurchaseOrder,
    KitchenPurchaseOrderItem,
    KitchenRecipe,
    KitchenRecipeIngredient,
)

weekly_procurement_bp = Blueprint("weekly_procurement", __name__)


def _as_date(raw: str | None) -> date:
    try:
        return date.fromisoformat(raw) if raw else date.today()
    except ValueError:
        return date.today()


def _week_bounds(day: date) -> tuple[date, date]:
    start = day - timedelta(days=day.weekday())
    return start, start + timedelta(days=6)


def _d(value) -> Decimal:
    return value if isinstance(value, Decimal) else Decimal(str(value or 0))


def _trim(value) -> str:
    text = format(_d(value), "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _expected_sources(week_start: date, week_end: date):
    """Rebuild ingredient demand from current school assignments + current BOM."""
    assignments = (
        KitchenMenuAssignment.query.join(KitchenMenuPlan)
        .filter(
            KitchenMenuPlan.service_date.between(week_start, week_end),
            KitchenMenuAssignment.service_status == "serving",
            KitchenMenuAssignment.headcount > 0,
        )
        .options(
            selectinload(KitchenMenuAssignment.school),
            selectinload(KitchenMenuAssignment.plan)
            .selectinload(KitchenMenuPlan.items)
            .selectinload(KitchenMenuPlanItem.recipe)
            .selectinload(KitchenRecipe.ingredients)
            .selectinload(KitchenRecipeIngredient.ingredient),
        )
        .all()
    )
    sources = defaultdict(lambda: defaultdict(lambda: {
        "required": Decimal("0"), "schools": [], "recipe_name": "", "recipe_id": None,
    }))
    totals = defaultdict(lambda: Decimal("0"))
    for assignment in assignments:
        service_date = assignment.plan.service_date
        school_name = assignment.school.name if assignment.school else f"學校 #{assignment.school_id}"
        for plan_item in assignment.plan.items:
            recipe = plan_item.recipe
            for component in recipe.ingredients:
                if component.quantity_status == "pending":
                    continue
                amount = _d(component.grams_per_person) * assignment.headcount
                if amount <= 0:
                    continue
                key = (service_date, component.ingredient_id)
                source_key = (recipe.id, recipe.name)
                row = sources[key][source_key]
                row["recipe_id"] = recipe.id
                row["recipe_name"] = recipe.name
                row["required"] += amount
                row["schools"].append({"name": school_name, "headcount": assignment.headcount})
                totals[key] += amount
    return sources, totals


def _orders_for_week(week_start: date, week_end: date):
    return (
        KitchenPurchaseOrder.query.filter(
            KitchenPurchaseOrder.service_date.between(week_start, week_end),
            KitchenPurchaseOrder.status != "cancelled",
        )
        .options(
            selectinload(KitchenPurchaseOrder.items).selectinload(KitchenPurchaseOrderItem.supplier),
            selectinload(KitchenPurchaseOrder.items).selectinload(KitchenPurchaseOrderItem.ingredient),
        )
        .order_by(KitchenPurchaseOrder.service_date, KitchenPurchaseOrder.id)
        .all()
    )


def _weekly_data(week_start: date, week_end: date):
    sources, expected_totals = _expected_sources(week_start, week_end)
    orders = _orders_for_week(week_start, week_end)
    item_keys = defaultdict(list)
    vendor_groups = defaultdict(lambda: defaultdict(list))
    anomalies = []

    for order in orders:
        for item in order.items:
            key = (order.service_date, item.ingredient_id)
            item_keys[key].append(item)
            expected = expected_totals.get(key, Decimal("0"))
            saved_required = _d(item.required_grams)
            if expected <= 0:
                anomalies.append({
                    "type": "orphan",
                    "date": order.service_date,
                    "ingredient": item.ingredient_name_snapshot,
                    "message": "採購單有此品項，但目前菜單／配方找不到需求來源。",
                })
            elif abs(expected - saved_required) > max(Decimal("0.01"), expected * Decimal("0.001")):
                anomalies.append({
                    "type": "mismatch",
                    "date": order.service_date,
                    "ingredient": item.ingredient_name_snapshot,
                    "message": f"目前菜單需求 {_trim(expected)}，採購單保存 {_trim(saved_required)}，數量不一致。",
                })

            source_rows = []
            day_total = expected_totals.get(key, Decimal("0"))
            for source in sources.get(key, {}).values():
                allocation = Decimal("0")
                if day_total > 0:
                    allocation = _d(item.actual_order_qty) * source["required"] / day_total
                source_rows.append({
                    **source,
                    "allocated_actual_qty": allocation,
                })
            source_rows.sort(key=lambda row: row["recipe_name"])
            supplier_name = (
                item.supplier_name_snapshot
                or (item.supplier.name if item.supplier else None)
                or "未指定廠商"
            )
            delivery_date = item.delivery_date or order.service_date
            vendor_groups[supplier_name][delivery_date].append({
                "item": item,
                "order": order,
                "sources": source_rows,
                "expected": expected,
            })

    for key, rows in item_keys.items():
        if len(rows) > 1:
            anomalies.append({
                "type": "duplicate",
                "date": key[0],
                "ingredient": rows[0].ingredient_name_snapshot,
                "message": f"同一天同一食材出現 {len(rows)} 筆採購資料，請確認是否重複。",
            })

    for key, expected in expected_totals.items():
        if expected > 0 and not item_keys.get(key):
            ingredient_id = key[1]
            sample = next(iter(sources[key].values()))
            recipe_names = "、".join(sorted({row["recipe_name"] for row in sources[key].values()}))
            ingredient_name = None
            for assignment_source in sources[key].values():
                # Name is not stored in the source bucket; resolve once by component id below if needed.
                break
            ingredient = db.session.get(__import__("models").KitchenIngredient, ingredient_id)
            ingredient_name = ingredient.name if ingredient else f"食材 #{ingredient_id}"
            anomalies.append({
                "type": "missing",
                "date": key[0],
                "ingredient": ingredient_name,
                "message": f"目前菜單需要此食材（來源：{recipe_names}），但採購單沒有對應品項。",
            })

    vendors = []
    for supplier_name in sorted(vendor_groups, key=str.casefold):
        deliveries = []
        for delivery_date in sorted(vendor_groups[supplier_name]):
            rows = vendor_groups[supplier_name][delivery_date]
            rows.sort(key=lambda row: row["item"].ingredient_name_snapshot.casefold())
            deliveries.append({"date": delivery_date, "rows": rows})
        vendors.append({"name": supplier_name, "deliveries": deliveries})

    anomalies.sort(key=lambda row: (row["date"], row["ingredient"].casefold(), row["type"]))
    return orders, vendors, anomalies


@weekly_procurement_bp.get("/weekly-procurement")
def weekly_procurement():
    selected = _as_date(request.args.get("week"))
    week_start, week_end = _week_bounds(selected)
    orders, vendors, anomalies = _weekly_data(week_start, week_end)
    total_items = sum(len(order.items) for order in orders)
    ordered_items = sum(1 for order in orders for item in order.items if item.ordered)
    return render_template(
        "kitchen/weekly_procurement.html",
        week_start=week_start,
        week_end=week_end,
        previous_week=week_start - timedelta(days=7),
        next_week=week_start + timedelta(days=7),
        vendors=vendors,
        anomalies=anomalies,
        total_items=total_items,
        ordered_items=ordered_items,
        trim_decimal=_trim,
    )


def _safe_sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", name or "未指定廠商")[:31] or "採購單"
    candidate = base
    index = 2
    while candidate in used:
        suffix = f"-{index}"
        candidate = (base[:31 - len(suffix)] + suffix)
        index += 1
    used.add(candidate)
    return candidate


@weekly_procurement_bp.get("/weekly-procurement.xlsx")
def weekly_procurement_export():
    selected = _as_date(request.args.get("week"))
    week_start, week_end = _week_bounds(selected)
    _orders, vendors, _anomalies = _weekly_data(week_start, week_end)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names = set()

    if not vendors:
        sheet = workbook.create_sheet("週採購單")
        sheet.append(["本週尚無採購資料"])
    for vendor in vendors:
        sheet = workbook.create_sheet(_safe_sheet_name(vendor["name"], used_names))
        sheet.append([vendor["name"]])
        sheet.append([f"週採購單：{week_start.isoformat()} ～ {week_end.isoformat()}"])
        sheet.append([])
        sheet.append(["交貨日期", "品項", "數量", "單位", "備註"])
        for cell in sheet[4]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        aggregated = defaultdict(lambda: {"qty": Decimal("0"), "unit": "", "notes": []})
        for delivery in vendor["deliveries"]:
            for row in delivery["rows"]:
                item = row["item"]
                key = (delivery["date"], item.ingredient_name_snapshot, item.purchase_unit_snapshot)
                aggregated[key]["qty"] += _d(item.actual_order_qty)
                aggregated[key]["unit"] = item.purchase_unit_snapshot
                if item.note and item.note not in aggregated[key]["notes"]:
                    aggregated[key]["notes"].append(item.note)
        for (delivery_date, ingredient_name, unit), data in sorted(aggregated.items(), key=lambda pair: (pair[0][0], pair[0][1].casefold())):
            sheet.append([
                delivery_date.isoformat(),
                ingredient_name,
                float(data["qty"]),
                unit,
                "；".join(data["notes"]),
            ])
        for column, width in {"A": 14, "B": 24, "C": 14, "D": 10, "E": 36}.items():
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"週採購單-{week_start.isoformat()}-{week_end.isoformat()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
