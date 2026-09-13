"""團膳共用相容鉤子：QR 練習入口與公版菜單 Excel。"""
from __future__ import annotations

from io import BytesIO
import re
from urllib.parse import quote

from flask import request

from blueprints.order_tool import order_bp

_PRACTICE_LINK = (
    '<a class="btn" href="/admin/order-tool/ai-menu/practice" '
    'target="_blank" rel="noopener">前往練習專區</a>'
)


def _safe_download_name(value: str) -> str:
    text = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", (value or "").strip())
    text = re.sub(r"\s+", " ", text).strip(" ._")
    if not text:
        return ""
    text = text[:80].rstrip(" ._")
    if not text.lower().endswith(".xlsx"):
        text += ".xlsx"
    return text


def _rewrite_public_menu_xlsx(response):
    """把 AI 公版匯出改成學校月菜單排列，並保持可匯回「總表」。"""
    matched = re.search(r"/ai-menu/drafts/(\d+)/public\.xlsx$", request.path)
    if not matched:
        return response

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from ai_models import KitchenMenuDraft
    from extensions import db
    from blueprints import ai_menu as ai_menu_module

    draft = db.session.get(KitchenMenuDraft, int(matched.group(1)))
    if draft is None:
        return response

    result = ai_menu_module._draft_result(draft)
    rows = ai_menu_module._view_rows(result, draft)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "公版菜單"

    sheet.merge_cells("A1:Q1")
    sheet["A1"] = f"{draft.name}　{draft.start_date.year}年{draft.start_date.month}月菜單"
    sheet["A1"].font = Font(size=18, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    headers = [
        "日期", "星期", "主食", "主菜", "副菜", "", "蔬菜", "湯品",
        "全穀雜糧類(份)", "豆魚蛋肉類(份)", "蔬菜類(份)",
        "油脂與堅果種子類(份)", "水果類(份)", "乳品類(份)",
        "熱量(大卡)", "營養狀態", "規則分數",
    ]
    sheet.append(headers)
    thin = Side(style="thin", color="666666")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="E8EEF5")
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def pick(row, category, slot=1):
        return next(
            (cell for cell in row["cells"] if cell["category"] == category and cell["slot_index"] == slot),
            None,
        )

    for row in rows:
        main_food = pick(row, "主食")
        main_dish = pick(row, "主菜")
        side_1 = pick(row, "副菜", 1)
        side_2 = pick(row, "副菜", 2)
        vegetable = pick(row, "青菜")
        soup = pick(row, "湯品")
        dessert = pick(row, "點心")
        if dessert and dessert.get("name"):
            if soup and soup.get("name"):
                soup = dict(soup)
                soup["name"] += f"／{dessert['name']}"
                if dessert.get("ingredients"):
                    soup["ingredients"] += "／" + dessert["ingredients"]
            else:
                soup = dessert
        dishes = [main_food, main_dish, side_1, side_2, vegetable, soup]
        report = result.days[row["service_date"]]
        servings = report.servings or {}
        status = "；".join(report.warnings or report.ok_messages) or "—"
        top = [row["service_date"], row["weekday"]]
        top += [(cell.get("name", "") if cell else "") for cell in dishes]
        top += [
            servings.get("grains") or "",
            servings.get("protein") or "",
            servings.get("vegetable") or "",
            servings.get("oil") or "",
            servings.get("fruit") or "",
            servings.get("dairy") or "",
            row["kcal"] if row["kcal"] is not None else "",
            status,
            round(result.penalty / 1000, 1),
        ]
        bottom = ["", ""]
        bottom += [(cell.get("ingredients", "") if cell else "") for cell in dishes]
        bottom += ["", "", "", "", "", "", "", "", ""]
        sheet.append(top)
        sheet.append(bottom)
        top_row = sheet.max_row - 1
        bottom_row = sheet.max_row
        for column in range(1, 18):
            for row_number in (top_row, bottom_row):
                cell = sheet.cell(row_number, column)
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in range(3, 9):
            sheet.cell(top_row, column).font = Font(size=14, bold=True)
            sheet.cell(bottom_row, column).font = Font(size=9, color="555555")
        for column in range(9, 18):
            sheet.cell(top_row, column).font = Font(size=9)
        sheet.row_dimensions[top_row].height = 34
        sheet.row_dimensions[bottom_row].height = 26

    widths = {
        "A": 12, "B": 7, "C": 21, "D": 21, "E": 19, "F": 19,
        "G": 17, "H": 21, "I": 10, "J": 10, "K": 9, "L": 11,
        "M": 8, "N": 8, "O": 11, "P": 30, "Q": 10,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "C3"
    sheet.sheet_view.showGridLines = False

    note = workbook.create_sheet("說明")
    note.append(["欄位", "說明"])
    note.append([
        "六大類份數",
        "依 Recipe BOM 每人用量與國民健康署食物代換表估算；油脂份數以總熱量扣除其他食物類標準熱量後反推。",
    ])
    note.append([
        "可匯回總表",
        "公版菜單保留日期、星期、主食、主菜、副菜、蔬菜、湯品欄位；既有總表匯入器會忽略第二列食材與右側份數欄。",
    ])
    note.append([
        "官方基準",
        "教育部「學校午餐食物內容及營養基準」109.12.28 修訂；份量定義依國民健康署食物代換表。",
    ])
    note.column_dimensions["A"].width = 20
    note.column_dimensions["B"].width = 90
    for row in note.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    body = output.getvalue()
    response.set_data(body)
    response.headers["Content-Length"] = str(len(body))
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return response


@order_bp.after_app_request
def _ai_menu_ui_compatibility(response):
    if request.path == "/punch/qrcode" and response.status_code == 200 and "text/html" in (response.content_type or "").lower():
        html = response.get_data(as_text=True)
        marker = "前往叫貨專區</a>"
        if marker in html and "前往練習專區</a>" not in html:
            html = html.replace(marker, marker + _PRACTICE_LINK, 1)
            response.set_data(html)

    if request.path.endswith("/public.xlsx") and response.status_code == 200:
        response = _rewrite_public_menu_xlsx(response)
        filename = _safe_download_name(request.args.get("filename", ""))
        if filename:
            response.headers["Content-Disposition"] = "attachment; filename*=UTF-8''" + quote(filename)
    return response
