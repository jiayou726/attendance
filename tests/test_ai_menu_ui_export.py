import json
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from ai_models import KitchenMealProfile, KitchenMenuDraft, KitchenMenuDraftItem
from app import create_app
from extensions import db
from models import KitchenIngredient, KitchenRecipe, KitchenRecipeIngredient


@pytest.fixture()
def app(tmp_path):
    db_path = tmp_path / "ai_menu_ui.db"
    app = create_app({
        "TESTING": True,
        "SECRET_KEY": "test-only-secret",
        "SQLALCHEMY_DATABASE_URI": f"sqlite:///{db_path}",
        "AUTO_CREATE_DB": True,
        "KITCHEN_CSRF_ENABLED": False,
        "ADMIN_MGR_PASSWORD": "test-password",
        "PRODUCTION": False,
    })
    yield app
    with app.app_context():
        db.session.remove()
        db.drop_all()


@pytest.fixture()
def client(app):
    return app.test_client()


def _seed_draft(app):
    with app.app_context():
        ingredient = KitchenIngredient(
            name="白米", base_unit="g", purchase_unit="kg",
            grams_per_purchase_unit=Decimal("1000"), unit_price=Decimal("0"),
            order_increment=Decimal("0.001"), active=True, kcal_per_100g=Decimal("354"),
        )
        first = KitchenRecipe(name="白米主菜", category="主菜", active=True)
        second = KitchenRecipe(name="糙米主菜", category="主菜", active=True)
        profile = KitchenMealProfile(code="test", name="測試", kcal_min=600, kcal_max=800, active=True)
        db.session.add_all([ingredient, first, second, profile]); db.session.flush()
        for recipe in (first, second):
            db.session.add(KitchenRecipeIngredient(
                recipe_id=recipe.id, ingredient_id=ingredient.id,
                grams_per_person=Decimal("50"), quantity_status="manual",
            ))
        draft = KitchenMenuDraft(
            name="測試公版菜單", profile_id=profile.id,
            start_date=date(2026, 6, 1), end_date=date(2026, 6, 1),
            include_weekends=False, meal_type="午餐",
            structure_json=json.dumps({"主菜": 1}),
            rules_json=json.dumps({
                "recipe_repeat_days": 0, "main_repeat_days": 0,
                "fish_per_week_min": 0, "fried_per_week_max": 1,
                "sweet_soup_per_week_max": 1,
                "exclude_incomplete_nutrition": True,
            }), status="draft",
        )
        db.session.add(draft); db.session.flush()
        db.session.add(KitchenMenuDraftItem(
            draft_id=draft.id, service_date=date(2026, 6, 1),
            category="主菜", slot_index=1, sort_order=11, recipe_id=first.id,
        ))
        db.session.commit()
        return draft.id, first.id


def test_ai_menu_uses_searchable_swap_and_only_daily_lock(app, client):
    draft_id, recipe_id = _seed_draft(app)

    result_page = client.get(f"/admin/order-tool/ai-menu/drafts/{draft_id}").get_data(as_text=True)
    assert f"/admin/order-tool/recipes/{recipe_id}" in result_page
    assert "/lock-day" in result_page
    assert "/lock-item" not in result_page

    swap_page = client.get(
        f"/admin/order-tool/ai-menu/drafts/{draft_id}/swap?date=2026-06-01&category=主菜&slot=1"
    ).get_data(as_text=True)
    assert 'class="dish-search-input"' in swap_page
    assert 'data-search-mode="select"' in swap_page
    assert 'id="recipe-search-data"' in swap_page
    assert '"id": 2' in swap_page
    assert "ai-swap-table" not in swap_page


def test_public_excel_matches_simple_two_row_menu_format(app, client):
    draft_id, _recipe_id = _seed_draft(app)
    response = client.get(f"/admin/order-tool/ai-menu/drafts/{draft_id}/public.xlsx")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), data_only=False)
    assert workbook.sheetnames == ["公版菜單"]
    sheet = workbook["公版菜單"]
    assert [sheet.cell(2, col).value for col in range(1, 17)] == [
        "日期", "星期", "主食", "主菜", "副菜", None, "蔬菜", "湯品",
        "全穀雜糧類(份)", "豆魚蛋肉類(份)", "蔬菜類(份)",
        "油脂與堅果種子類(份)", "水果類", "乳品類", "熱量(大卡)", "三章1Q",
    ]
    assert sheet["D3"].value == "白米主菜"
    assert sheet["D4"].value == "白米"
    assert "50" not in sheet["D4"].value
    assert sheet["P3"].value == "✓"
    assert "狀態" not in [cell.value for cell in sheet[2]]
    assert "規則分數" not in [cell.value for cell in sheet[2]]
