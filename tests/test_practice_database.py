from io import BytesIO
import os

import pytest

from openpyxl import Workbook

from datetime import datetime, timedelta, timezone

from sqlalchemy import func, select, text

from app import create_app
from extensions import db
from models import KitchenMenuPlan, KitchenRecipe, KitchenSchool
from services.practice_database import cleanup_expired_accounts, list_accounts, workspace_engine


def _app(tmp_path, *, limit=20, ttl_days=10):
    main_path = tmp_path / "formal.db"
    practice_path = tmp_path / "practice-control.db"
    return create_app({
        "TESTING": True,
        "SECRET_KEY": "practice-test-secret",
        "SQLALCHEMY_DATABASE_URI": f"sqlite:///{main_path}",
        "SQLALCHEMY_BINDS": {"practice": f"sqlite:///{practice_path}"},
        "AUTO_CREATE_DB": True,
        "KITCHEN_CSRF_ENABLED": False,
        "PRACTICE_ACCOUNT_LIMIT": limit,
        "PRACTICE_ACCOUNT_TTL_DAYS": ttl_days,
        "PRACTICE_ACCOUNTS_ENABLED": True,
        "PRODUCTION": False,
    })


def _login(client, name):
    response = client.post(
        "/admin/order-tool/ai-menu/practice/login",
        data={"name": name},
    )
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/admin/order-tool/")


def _create_plan(client, name="練習週菜單"):
    return client.post("/admin/order-tool/plans", data={
        "service_date": "2026-10-05",
        "meal_type": "午餐",
        "name": name,
    })


def test_named_accounts_resume_across_computers_and_remain_isolated(tmp_path):
    app = _app(tmp_path)
    with app.app_context():
        db.session.add(KitchenSchool(name="正式學校"))
        db.session.commit()

    first_computer = app.test_client()
    _login(first_computer, "王小明")
    assert _create_plan(first_computer).status_code == 302

    second_computer = app.test_client()
    _login(second_computer, "王小明")
    response = second_computer.get("/admin/order-tool/summary?week=2026-10-05")
    assert response.status_code == 200
    assert "練習週菜單" in response.get_data(as_text=True)

    different_person = app.test_client()
    _login(different_person, "李小華")
    response = different_person.get("/admin/order-tool/summary?week=2026-10-05")
    assert response.status_code == 200
    assert "練習週菜單" not in response.get_data(as_text=True)

    with app.app_context():
        assert KitchenSchool.query.filter_by(name="正式學校").one()
        assert KitchenMenuPlan.query.count() == 0
        assert len(list_accounts()) == 2


def test_practice_master_data_is_editable_and_isolated(tmp_path):
    app = _app(tmp_path)
    with app.app_context():
        db.session.add(KitchenSchool(name="正式學校"))
        db.session.commit()
    client = app.test_client()
    _login(client, "練習者")

    response = client.post("/admin/order-tool/schools", data={"name": "練習新增學校"})
    assert response.status_code == 302
    assert "練習新增學校" in client.get("/admin/order-tool/schools").get_data(as_text=True)
    with app.app_context():
        assert KitchenSchool.query.filter_by(name="練習新增學校").first() is None


def test_capacity_and_expiry_release_slots(tmp_path):
    app = _app(tmp_path, limit=2)
    one = app.test_client()
    two = app.test_client()
    three = app.test_client()
    _login(one, "一號")
    _login(two, "二號")

    response = three.post("/admin/order-tool/ai-menu/practice/login", data={"name": "三號"})
    assert response.status_code == 200
    assert "已達 2 人上限" in response.get_data(as_text=True)

    with app.app_context():
        engine = db.engines["practice"]
        stale = datetime.now(timezone.utc) - timedelta(days=11)
        with engine.begin() as connection:
            connection.execute(text(
                "UPDATE practice_account SET last_active_at=:stale WHERE normalized_name=:name"
            ), {"stale": stale, "name": "一號"})
        assert cleanup_expired_accounts() == 1

    _login(three, "三號")
    with app.app_context():
        assert {row["display_name"] for row in list_accounts()} == {"二號", "三號"}


def test_admin_link_can_reset_and_delete_account(tmp_path):
    app = _app(tmp_path)
    learner = app.test_client()
    _login(learner, "管理測試")
    _create_plan(learner)

    admin = app.test_client()
    with admin.session_transaction() as admin_session:
        admin_session["role"] = "hr"
        admin_session["_practice_admin_csrf"] = "admin-test-token"

    employee_page = admin.get("/admin/").get_data(as_text=True)
    assert "出勤卡查詢" in employee_page
    assert "練習人員管理" in employee_page
    management_page = admin.get("/admin/practice-users/")
    assert management_page.status_code == 200
    assert "管理測試" in management_page.get_data(as_text=True)

    with app.app_context():
        account = list_accounts()[0]
        account_id = account["id"]
        assert account["plan_count"] == 1
    response = admin.post(
        f"/admin/practice-users/{account_id}/reset",
        data={"_csrf_token": "admin-test-token"},
    )
    assert response.status_code == 302
    with app.app_context():
        assert list_accounts()[0]["plan_count"] == 0

    response = admin.post(
        f"/admin/practice-users/{account_id}/delete",
        data={"_csrf_token": "admin-test-token"},
    )
    assert response.status_code == 302
    with app.app_context():
        assert list_accounts() == []


def test_practice_excel_import_new_recipe_stays_private(tmp_path):
    """Excel imports may introduce dishes absent from the formal recipe catalog."""
    app = _app(tmp_path)
    with app.app_context():
        db.session.add(KitchenRecipe(name="正式既有菜", category="主菜"))
        db.session.commit()

    client = app.test_client()
    _login(client, "菜單匯入練習者")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "一般菜單"
    sheet.append(["日期", "星期", "主菜", "湯品"])
    sheet.append(["2026/10/22", "四", "正式既有菜", "練習新湯品"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    imported = client.post(
        "/admin/order-tool/summary/import",
        data={"menu_file": (output, "2026年10月練習菜單.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert imported.status_code == 200
    assert "匯入完成" in imported.get_data(as_text=True)
    assert "練習新湯品" in imported.get_data(as_text=True)

    with app.app_context():
        # The formal master catalog has not been modified by a practice user.
        assert KitchenRecipe.query.filter_by(name="練習新湯品").count() == 0
        assert KitchenRecipe.query.filter_by(name="正式既有菜").count() == 1
        engine = workspace_engine(list_accounts()[0])
        with engine.connect() as connection:
            private_recipes = connection.execute(
                select(KitchenRecipe.__table__)
            ).mappings().all()
            private_plans = connection.execute(
                select(KitchenMenuPlan.__table__)
            ).mappings().all()
        assert {row["name"] for row in private_recipes} == {"正式既有菜", "練習新湯品"}
        assert len(private_plans) == 1

    # The same account can safely import the same workbook again: no duplicates.
    repeated = client.post(
        "/admin/order-tool/summary/import",
        data={"menu_file": (_practice_menu_file(), "2026年10月練習菜單.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert repeated.status_code == 200
    assert "略過" in repeated.get_data(as_text=True)
    with app.app_context():
        engine = workspace_engine(list_accounts()[0])
        with engine.connect() as connection:
            assert len(connection.execute(select(KitchenRecipe.__table__)).all()) == 2


def _practice_menu_file():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "一般菜單"
    sheet.append(["日期", "星期", "主菜", "湯品"])
    sheet.append(["2026/10/22", "四", "正式既有菜", "練習新湯品"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _realistic_november_menu():
    """Match the uploaded 中平 November template: offset date column,
    six dish columns, and a separate ingredient-description row per day.
    """
    from datetime import date

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "中平115.11月"
    sheet.append([None] * 16)
    sheet.append([None, "日期", "星期", "主食", "主菜", "副菜", "副菜", "青菜", "湯品",
                  "全穀雜糧類(份)", "豆魚蛋肉類(份)", "蔬菜類(份)", "油脂類(份)",
                  "其他類(份)", "熱量(大卡)", "三章1Q"])
    for day in range(1, 31):
        d = date(2026, 11, day)
        if d.weekday() >= 5:
            continue
        sheet.append([
            None, datetime(2026, 11, day), "一" if d.weekday() == 0 else "大塊肉",
            f"飯{day}", f"主菜{day}", f"副菜甲{day}", f"副菜乙{day}",
            f"青菜{day}", f"湯{day}", 5.6, 2.2, 2, 2.2, None, 706, "V",
        ])
        sheet.append([None, None, None, None, "豬肉Q.洋蔥Q/燒",
                      "豆腐.菇Q/煮", "玉米粒S/煮", None, "雞蛋Q"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_practice_import_month_template_is_complete_and_idempotent(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    _login(client, "完整月份匯入")
    response = client.post(
        "/admin/order-tool/summary/import",
        data={"menu_file": (_realistic_november_menu(), "中平11月菜單.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "匯入完成" in page
    assert "21 天新增 126 道菜" in page

    with app.app_context():
        assert KitchenMenuPlan.query.count() == 0  # Formal database stays untouched.
        assert KitchenRecipe.query.count() == 0
        engine = workspace_engine(list_accounts()[0])
        with engine.connect() as connection:
            assert connection.execute(select(func.count()).select_from(
                KitchenMenuPlan.__table__
            )).scalar_one() == 21
            assert connection.execute(select(func.count()).select_from(
                KitchenRecipe.__table__
            )).scalar_one() == 126

    again = client.post(
        "/admin/order-tool/summary/import",
        data={"menu_file": (_realistic_november_menu(), "中平11月菜單.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "0 天新增 0 道菜" in again.get_data(as_text=True)


@pytest.mark.skipif(not os.environ.get("PRACTICE_POSTGRES_TEST_URL"),
                    reason="Requires dedicated disposable PostgreSQL for integration tests")
def test_practice_postgres_excel_import_uses_private_master_copy():
    # Never point this test at a shared or production database.
    postgres_url = os.environ["PRACTICE_POSTGRES_TEST_URL"]
    app = create_app({
        "TESTING": True,
        "SECRET_KEY": "postgres-practice-test-secret",
        "SQLALCHEMY_DATABASE_URI": postgres_url,
        "SQLALCHEMY_BINDS": {"practice": postgres_url},
        "AUTO_CREATE_DB": True,
        "KITCHEN_CSRF_ENABLED": False,
        "PRACTICE_ACCOUNT_LIMIT": 20,
        "PRACTICE_ACCOUNT_TTL_DAYS": 10,
        "PRACTICE_ACCOUNTS_ENABLED": True,
        "PRODUCTION": False,
    })
    with app.app_context():
        db.session.add(KitchenRecipe(name="正式既有菜", category="主菜"))
        db.session.commit()
    client = app.test_client()
    _login(client, "PostgreSQL 練習匯入")
    response = client.post(
        "/admin/order-tool/summary/import",
        data={"menu_file": (_realistic_november_menu(), "中平11月菜單.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "21 天新增 126 道菜" in page

    with app.app_context():
        assert KitchenRecipe.query.count() == 1  # No formal recipes created.
        assert KitchenMenuPlan.query.count() == 0
        account = list_accounts()[0]
        engine = workspace_engine(account)
        with engine.connect() as connection:
            private_recipes = connection.execute(select(KitchenRecipe.__table__)).mappings().all()
            assert len(private_recipes) == 127
            assert sum(row["id"] > 0 for row in private_recipes) == 127
            assert connection.execute(select(func.count()).select_from(
                KitchenMenuPlan.__table__
            )).scalar_one() == 21
            assert connection.execute(select(func.count()).select_from(
                db.metadata.tables["kitchen_menu_plan_item"]
            )).scalar_one() == 126
            plan_ids = connection.execute(select(
                db.metadata.tables["kitchen_menu_plan_item"].c.plan_id
            )).scalars().all()
            assert len(set(plan_ids)) == 21, {
                plan_id: plan_ids.count(plan_id) for plan_id in set(plan_ids)
            }
            plan = KitchenMenuPlan.__table__
            item = db.metadata.tables["kitchen_menu_plan_item"]
            recipe = KitchenRecipe.__table__
            links = connection.execute(
                select(plan.c.service_date, recipe.c.name)
                .select_from(plan.join(item, plan.c.id == item.c.plan_id)
                                 .join(recipe, recipe.c.id == item.c.recipe_id))
            ).all()
            mismatched = [(str(day), name) for day, name in links
                          if not name.endswith(str(day.day))]
            assert not mismatched, mismatched[:30]

    repeated = client.post(
        "/admin/order-tool/summary/import",
        data={"menu_file": (_realistic_november_menu(), "中平11月菜單.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    repeated_page = repeated.get_data(as_text=True)
    assert "0 天新增 0 道菜" in repeated_page, repeated_page.partition('class="flash')[2][:1000]


def test_practice_week_survives_relogin_and_weekly_report_is_reachable(tmp_path):
    app = _app(tmp_path, ttl_days=0)
    learner = app.test_client()
    _login(learner, "週採購練習人員")
    assert _create_plan(learner, "跨登入保留菜單").status_code == 302
    page = learner.get("/admin/order-tool/summary?week=2026-10-05")
    assert "跨登入保留菜單" in page.get_data(as_text=True)

    home = learner.get("/admin/order-tool/").get_data(as_text=True)
    assert "查看練習週採購單" in home
    assert "week=2026-10-05" in home
    weekly = learner.get("/admin/order-tool/weekly-procurement/summary.json")
    assert weekly.status_code == 200
    assert weekly.get_json()["week_start"] == "2026-10-05"
    assert "2026/10/05" in learner.get(
        "/admin/order-tool/weekly-procurement"
    ).get_data(as_text=True)

    learner.get("/admin/order-tool/ai-menu/practice/logout")
    _login(learner, "週採購練習人員")
    assert "跨登入保留菜單" in learner.get(
        "/admin/order-tool/summary"
    ).get_data(as_text=True)
    assert learner.get(
        "/admin/order-tool/weekly-procurement/summary.json"
    ).get_json()["week_start"] == "2026-10-05"

    other = app.test_client()
    _login(other, "另一位練習人員")
    assert "跨登入保留菜單" not in other.get(
        "/admin/order-tool/summary?week=2026-10-05"
    ).get_data(as_text=True)

    with app.app_context():
        engine = db.engines["practice"]
        stale = datetime.now(timezone.utc) - timedelta(days=90)
        with engine.begin() as connection:
            connection.execute(text(
                "UPDATE practice_account SET last_active_at=:stale"
            ), {"stale": stale})
        assert cleanup_expired_accounts() == 0
        assert len(list_accounts()) == 2
        assert all(account["days_remaining"] is None for account in list_accounts())


def test_practice_explicit_week_changes_the_weekly_report(tmp_path):
    app = _app(tmp_path)
    learner = app.test_client()
    _login(learner, "週次切換練習")
    _create_plan(learner, "舊週資料")
    learner.get("/admin/order-tool/summary?week=2026-11-02")
    assert learner.get(
        "/admin/order-tool/weekly-procurement/summary.json"
    ).get_json()["week_start"] == "2026-11-02"
    assert "舊週資料" in learner.get(
        "/admin/order-tool/summary?week=2026-10-05"
    ).get_data(as_text=True)


def test_existing_sqlite_practice_workspace_is_auto_upgraded_without_data_loss(tmp_path):
    from services.practice_database import workspace_engine
    app = _app(tmp_path)
    client = app.test_client()
    _login(client, "舊資料升級練習")
    _create_plan(client, "升級前保留菜單")
    with app.app_context():
        account = list_accounts()[0]
        engine = workspace_engine(account)
        with engine.begin() as connection:
            connection.execute(text(
                "ALTER TABLE kitchen_menu_plan_item DROP COLUMN ingredients_override_json"
            ))
            connection.execute(text(
                "ALTER TABLE kitchen_daily_dish_note DROP COLUMN count_override_fields"
            ))
        app.extensions["practice_estimate_schema_checked"].discard(
            account["workspace_key"]
        )
    response = client.get("/admin/order-tool/summary?week=2026-10-05")
    assert response.status_code == 200
    assert "升級前保留菜單" in response.get_data(as_text=True)
    with app.app_context():
        engine = workspace_engine(list_accounts()[0])
        from sqlalchemy import inspect
        assert "ingredients_override_json" in {
            x["name"] for x in inspect(engine).get_columns("kitchen_menu_plan_item")
        }
        assert "count_override_fields" in {
            x["name"] for x in inspect(engine).get_columns("kitchen_daily_dish_note")
        }


@pytest.mark.skipif(not os.environ.get("PRACTICE_POSTGRES_TEST_URL"),
                    reason="Requires dedicated disposable PostgreSQL")
def test_practice_postgres_legacy_view_schema_upgrade_keeps_menu_and_private_recipe():
    """Old PostgreSQL accounts must become writable without losing plans or IDs."""
    from services.practice_database import MASTER_TABLES, workspace_engine
    from sqlalchemy import inspect

    url = os.environ["PRACTICE_POSTGRES_TEST_URL"]
    app = create_app({
        "TESTING": True,
        "SECRET_KEY": "legacy-practice-test-secret",
        "SQLALCHEMY_DATABASE_URI": url,
        "SQLALCHEMY_BINDS": {"practice": url},
        "AUTO_CREATE_DB": True,
        "KITCHEN_CSRF_ENABLED": False,
        "PRACTICE_ACCOUNTS_ENABLED": True,
        "PRODUCTION": False,
    })
    with app.app_context():
        db.session.add(KitchenRecipe(name="舊帳號同名菜", category="主菜"))
        db.session.commit()
    client = app.test_client()
    _login(client, "Postgres 舊帳號保留測試")
    _create_plan(client, "遷移前的菜單")
    with app.app_context():
        account = next(x for x in list_accounts()
                       if x["display_name"] == "Postgres 舊帳號保留測試")
        key = account["workspace_key"]
        db.session.remove()
        with db.engines["practice"].begin() as connection:
            for table in MASTER_TABLES:
                connection.execute(text(f'DROP TABLE "{key}"."{table}"'))
                connection.execute(text(
                    f'CREATE VIEW "{key}"."{table}" '
                    f'WITH (security_invoker=true) AS '
                    f'SELECT * FROM public."{table}" OFFSET 0'
                ))
            connection.execute(text(
                f'CREATE TABLE "{key}".practice_recipe_overlay '
                '(LIKE public.kitchen_recipe INCLUDING CONSTRAINTS)'
            ))
            connection.execute(text(
                f'INSERT INTO "{key}".practice_recipe_overlay '
                '(id, name, category, active, created_at, updated_at) '
                "VALUES (-11, '舊帳號同名菜', '主菜', false, now(), now())"
            ))
            connection.execute(text(
                f'INSERT INTO "{key}".kitchen_menu_plan_item '
                '(plan_id, recipe_id, sort_order) '
                f'SELECT id, -11, 1 FROM "{key}".kitchen_menu_plan LIMIT 1'
            ))
            connection.execute(text(
                f'ALTER TABLE "{key}".kitchen_menu_plan_item '
                'DROP COLUMN ingredients_override_json'
            ))
            connection.execute(text(
                f'ALTER TABLE "{key}".kitchen_daily_dish_note '
                'DROP COLUMN count_override_fields'
            ))
        app.extensions["practice_estimate_schema_checked"].discard(key)

    response = client.get("/admin/order-tool/summary?week=2026-10-05")
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "遷移前的菜單" in page
    assert "練習原有-11" in page
    with app.app_context():
        account = next(x for x in list_accounts()
                       if x["display_name"] == "Postgres 舊帳號保留測試")
        engine = workspace_engine(account)
        with engine.connect() as connection:
            assert connection.execute(select(func.count()).select_from(
                KitchenMenuPlan.__table__
            )).scalar_one() == 1
            assert connection.execute(select(func.count()).select_from(
                db.metadata.tables["kitchen_menu_plan_item"]
            )).scalar_one() == 1
            assert connection.execute(text(
                f'SELECT count(*) FROM "{key}".kitchen_recipe WHERE id=-11'
            )).scalar_one() == 1
            assert "ingredients_override_json" in {
                c["name"] for c in inspect(connection).get_columns(
                    "kitchen_menu_plan_item", schema=key
                )
            }
            assert "count_override_fields" in {
                c["name"] for c in inspect(connection).get_columns(
                    "kitchen_daily_dish_note", schema=key
                )
            }
            assert connection.execute(text(
                'SELECT relkind FROM pg_class c JOIN pg_namespace n '
                'ON n.oid=c.relnamespace WHERE n.nspname=:key '
                "AND c.relname='kitchen_recipe'"
            ), {"key": key}).scalar_one() == "r"
        assert KitchenRecipe.query.filter_by(name="舊帳號同名菜").count() == 1


@pytest.mark.skipif(not os.environ.get("PRACTICE_POSTGRES_TEST_URL"),
                    reason="Requires dedicated disposable PostgreSQL")
def test_practice_postgres_master_copy_handles_public_column_order_and_extra_fields():
    """Prevent INSERT ... SELECT * regression on migrated Supabase schemas."""
    from models import KitchenSchool
    url = os.environ["PRACTICE_POSTGRES_TEST_URL"]
    app = create_app({
        "TESTING": True,
        "SECRET_KEY": "column-order-practice-test",
        "SQLALCHEMY_DATABASE_URI": url,
        "SQLALCHEMY_BINDS": {"practice": url},
        "AUTO_CREATE_DB": True,
        "KITCHEN_CSRF_ENABLED": False,
        "PRACTICE_ACCOUNTS_ENABLED": True,
        "PRODUCTION": False,
    })
    with app.app_context():
        db.session.add(KitchenSchool(
            name="欄位順序驗證學校", default_headcount=101,
            default_vegetarian_headcount=7, active=True,
        ))
        db.session.commit()
        # Simulate an additive formal migration that a new practice table does
        # not have. SELECT * would now fail, explicit named columns must work.
        with db.engine.begin() as conn:
            conn.execute(text(
                "ALTER TABLE public.kitchen_school "
                "ADD COLUMN IF NOT EXISTS practice_copy_test_extra BOOLEAN DEFAULT false"
            ))
    try:
        client = app.test_client()
        _login(client, "Postgres 欄位順序防呆")
        result = client.get("/admin/order-tool/schools")
        assert result.status_code == 200
        assert "欄位順序驗證學校" in result.get_data(as_text=True)
        with app.app_context():
            from services.practice_database import list_accounts, workspace_engine
            from sqlalchemy import select
            account = next(a for a in list_accounts() if a["display_name"] == "Postgres 欄位順序防呆")
            with workspace_engine(account).connect() as conn:
                school = conn.execute(select(KitchenSchool.__table__)
                    .where(KitchenSchool.__table__.c.name == "欄位順序驗證學校")).mappings().one()
                assert school["default_headcount"] == 101
                assert school["default_vegetarian_headcount"] == 7
    finally:
        with app.app_context():
            with db.engine.begin() as conn:
                conn.execute(text(
                    "ALTER TABLE public.kitchen_school "
                    "DROP COLUMN IF EXISTS practice_copy_test_extra"
                ))


def test_practice_session_loss_cannot_fall_back_to_formal_menu(tmp_path):
    """Practice must not accidentally display/write formal data if its cookie session expires."""
    app = _app(tmp_path)
    with app.app_context():
        db.session.add(KitchenMenuPlan(
            service_date=datetime(2026, 11, 2).date(),
            meal_type="午餐", name="正式區專用機密菜單"))
        db.session.commit()
    browser = app.test_client()
    _login(browser, "隔離狀態驗證")
    assert browser.get_cookie("kitchen_practice_mode", path="/admin/order-tool").value == "1"
    content = browser.get("/admin/order-tool/summary?week=2026-11-02").get_data(as_text=True)
    assert "正式區專用機密菜單" not in content

    # Simulate loss of the signed Flask session after a deploy while the
    # persistent practice-browser marker remains intact.
    with browser.session_transaction() as sess:
        sess.clear()
    get_result = browser.get("/admin/order-tool/summary?week=2026-11-02")
    assert get_result.status_code == 302
    assert "/ai-menu/practice/login" in get_result.headers["Location"]
    post_result = browser.post("/admin/order-tool/plans", data={
        "service_date": "2026-11-02", "meal_type": "午餐",
        "name": "絕不能進入正式區的練習菜單",
    })
    assert post_result.status_code == 302
    assert "/ai-menu/practice/login" in post_result.headers["Location"]
    with app.app_context():
        assert KitchenMenuPlan.query.filter_by(
            name="絕不能進入正式區的練習菜單").count() == 0

    # The explicit Leave Practice action resets only the browser mode.
    _login(browser, "隔離狀態驗證")
    browser.get("/admin/order-tool/ai-menu/practice/logout")
    assert browser.get_cookie("kitchen_practice_mode", path="/admin/order-tool") is None
    content = browser.get("/admin/order-tool/summary?week=2026-11-02").get_data(as_text=True)
    assert "正式區專用機密菜單" in content


def test_old_practice_sessions_gain_mode_cookie_before_session_expiry(tmp_path):
    """An existing signed practice session gets the fail-closed mode marker."""
    app = _app(tmp_path)
    browser = app.test_client()
    _login(browser, "保護既有練習登入")
    browser.delete_cookie("kitchen_practice_mode", path="/admin/order-tool")
    assert browser.get_cookie("kitchen_practice_mode", path="/admin/order-tool") is None
    assert browser.get("/admin/order-tool/summary").status_code == 200
    assert browser.get_cookie("kitchen_practice_mode", path="/admin/order-tool").value == "1"


def test_practice_login_page_has_no_formal_navigation(tmp_path):
    app = _app(tmp_path)
    browser = app.test_client()
    html = browser.get("/admin/order-tool/ai-menu/practice/login").get_data(as_text=True)
    assert "進入練習帳號" in html
    assert 'aria-label="團膳管理主選單"' not in html


def test_formal_and_practice_menu_and_master_edits_never_synchronize(tmp_path):
    """Both directions are isolated, including unchanged formal master snapshots."""
    app = _app(tmp_path)
    with app.app_context():
        db.session.add(KitchenSchool(name="正式原始學校", default_headcount=99))
        db.session.add(KitchenMenuPlan(
            service_date=datetime(2026, 11, 2).date(),
            meal_type="午餐", name="正式菜單不可顯示於練習區"))
        db.session.commit()

    learner = app.test_client()
    _login(learner, "獨立新帳號練習者")
    practice_page = learner.get("/admin/order-tool/summary?week=2026-11-02")
    assert practice_page.status_code == 200
    assert "正式菜單不可顯示於練習區" not in practice_page.get_data(as_text=True)
    assert "正式原始學校" in learner.get(
        "/admin/order-tool/schools"
    ).get_data(as_text=True)

    # Modify the formal master *after* the practice account was created.
    with app.app_context():
        school = KitchenSchool.query.filter_by(name="正式原始學校").one()
        school.name = "正式後來改名學校"
        db.session.commit()
    practice_schools = learner.get(
        "/admin/order-tool/schools"
    ).get_data(as_text=True)
    assert "正式原始學校" in practice_schools
    assert "正式後來改名學校" not in practice_schools

    learner.post("/admin/order-tool/plans", data={
        "service_date": "2026-11-02", "meal_type": "午餐",
        "name": "只存在練習區的菜單",
    })
    assert "只存在練習區的菜單" in learner.get(
        "/admin/order-tool/summary?week=2026-11-02"
    ).get_data(as_text=True)

    # Formal data is still unchanged and the practice plan is not visible.
    formal = app.test_client()
    formal_page = formal.get(
        "/admin/order-tool/summary?week=2026-11-02"
    ).get_data(as_text=True)
    assert "正式菜單不可顯示於練習區" in formal_page
    assert "只存在練習區的菜單" not in formal_page
    with app.app_context():
        assert KitchenMenuPlan.query.filter_by(
            name="只存在練習區的菜單").count() == 0
        assert KitchenSchool.query.filter_by(
            name="正式原始學校").count() == 0
